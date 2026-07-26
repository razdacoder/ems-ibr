"""DOCX and CSV renderers for the Hall Directory / VISA documents.

Parallel to :func:`ems.pdf_gen.build_pdf` — all three consume the same
``ems.directory.build_payload`` output so the formats stay in step.

DOCX mirrors the PDF (institution branding, one page per slot, same table).
CSV is a flat data export instead: one row per record with the slot's date and
period on every row, so several slots concatenate into something sortable and
filterable. Branding is deliberately omitted there — it would corrupt the
header row for anything consuming the file as data.
"""

import csv
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from ems.branding import add_document_branding

_HALL_HEADERS = ["Hall", "Class Name", "No. Of Students", "Matric Numbers"]
_HALL_CSV_HEADERS = ["Date", "Period", *_HALL_HEADERS]
_VISA_CSV_HEADERS = ["Date", "Period", "Department", "Department Name", "Classes"]


def _add_title(doc, title: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(12)


def _hall_slot_docx(doc, slot: dict) -> None:
    _add_title(doc, slot["title"])
    rows = slot.get("rows") or []
    if not rows:
        doc.add_paragraph("No seat allocation generated for this slot.")
        return
    table = doc.add_table(rows=1, cols=len(_HALL_HEADERS))
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells, _HALL_HEADERS):
        run = cell.paragraphs[0].add_run(heading)
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        cells[0].text = row["hall"]
        cells[1].text = row["class_name"]
        cells[2].text = str(row["count"])
        cells[3].text = row["matric_range"]


def _visa_slot_docx(doc, slot: dict) -> None:
    _add_title(doc, slot["title"])
    codes = slot.get("codes") or []
    text = ", ".join(codes) if codes else "No classes scheduled for this slot."
    # Single-cell bordered table, matching the boxed VISA block in the PDF.
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = text


def build_docx(doc_type: str, payload: list[dict], settings_obj) -> bytes:
    """Render the payload to a multi-page DOCX (one page per slot)."""
    document = Document()
    if not payload:
        add_document_branding(document, settings_obj)
        document.add_paragraph("No data for the selected scope.")
    for i, slot in enumerate(payload):
        if i > 0:
            document.add_page_break()
        add_document_branding(document, settings_obj)
        if doc_type == "hall":
            _hall_slot_docx(document, slot)
        else:
            _visa_slot_docx(document, slot)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def build_csv(doc_type: str, payload: list[dict]) -> bytes:
    """Render the payload to a flat CSV covering every slot in scope."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    if doc_type == "hall":
        writer.writerow(_HALL_CSV_HEADERS)
        for slot in payload:
            for row in slot.get("rows") or []:
                writer.writerow([
                    slot["date"],
                    slot["period"],
                    row["hall"],
                    row["class_name"],
                    row["count"],
                    row["matric_range"],
                ])
    else:
        writer.writerow(_VISA_CSV_HEADERS)
        for slot in payload:
            for group in slot.get("groups") or []:
                writer.writerow([
                    slot["date"],
                    slot["period"],
                    group["department"],
                    group["department_name"],
                    " ".join(group["codes"]),
                ])
    # utf-8-sig so Excel picks up the encoding and renders any non-ASCII
    # institution or class names correctly on open.
    return buf.getvalue().encode("utf-8-sig")
