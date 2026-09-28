"""Attendance sheets: one DOCX per (course, hall, class)."""
import io
import zipfile

from django.test import RequestFactory, TestCase
from django.urls import reverse
from docx import Document
from rest_framework.test import APIClient

from ems.csv_gen import export_arrangements
from ems.models import SeatArrangement, User

from .helpers import enrol, make_class, make_course, make_department, make_hall

DATE, PERIOD = "2026-08-24", "AM"


def read_sheets(response):
    """zip response -> {filename: {"text": all paragraphs, "matrics": [...], "seats": [...]}}"""
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    sheets = {}
    for name in archive.namelist():
        doc = Document(io.BytesIO(archive.read(name)))
        rows = [r for r in doc.tables[0].rows[1:] if r.cells[1].text]
        sheets[name] = {
            "text": "\n".join(p.text for p in doc.paragraphs),
            "matrics": [r.cells[1].text for r in rows],
            "seats": [int(r.cells[3].text) for r in rows],
        }
    return sheets


class AttendanceFixture(TestCase):
    """Two classes from different departments sit the same course in one hall."""

    def setUp(self):
        self.hall = make_hall("BB 1", 8, 10)
        self.course = make_course("GNS 101")
        self.accounting = make_class(3, make_department("ac"), name="HND I")
        self.banking = make_class(3, make_department("bf"), name="HND I")
        # Interleaved matrics: sorting on matric alone would mix the classes.
        self.seat(self.accounting, ["2400000001", "2400000003", "2400000005"], first_seat=1)
        self.seat(self.banking, ["2400000002", "2400000004", "2400000006"], first_seat=41)

    def seat(self, cls, matrics, first_seat):
        for n, student in enumerate(enrol(cls, matrics)):
            SeatArrangement.objects.create(
                date=DATE, period=PERIOD, student=student, seat_number=first_seat + 2 * n,
                hall=self.hall, course=self.course, cls=cls,
            )


class BulkAttendanceExportTests(AttendanceFixture):
    def export(self):
        request = RequestFactory().get("/", {"date": DATE, "period": PERIOD})
        return read_sheets(export_arrangements(request))

    def test_gives_each_class_its_own_sheet(self):
        # Regression: sheets were keyed on (course, hall), so both classes
        # landed on one sheet labelled with only the first class.
        sheets = self.export()

        self.assertEqual(len(sheets), 2)
        self.assertEqual(
            sorted(s["matrics"][0] for s in sheets.values()), ["2400000001", "2400000002"]
        )

    def test_each_sheet_lists_only_its_own_class_under_its_own_label(self):
        sheets = self.export()

        for sheet in sheets.values():
            label = "AC HND I" if sheet["matrics"][0].endswith("1") else "BF HND I"
            own = {"AC HND I": ["2400000001", "2400000003", "2400000005"],
                   "BF HND I": ["2400000002", "2400000004", "2400000006"]}[label]
            self.assertIn(f"LEVEL/CLASS: {label}", sheet["text"])
            self.assertEqual(sheet["matrics"], own)

    def test_seat_numbers_rise_with_matric_numbers_on_every_sheet(self):
        for sheet in self.export().values():
            self.assertEqual(sheet["seats"], sorted(sheet["seats"]))

    def test_sheet_file_names_are_unique_and_name_the_class(self):
        names = list(self.export())

        self.assertEqual(len(set(names)), 2)
        self.assertTrue(any("AC HND I" in n for n in names), names)
        self.assertTrue(any("BF HND I" in n for n in names), names)

    def test_unplaced_students_are_left_off_the_sheets(self):
        SeatArrangement.objects.filter(student__matric_no="2400000005").update(seat_number=None)

        sheets = self.export()

        all_matrics = [m for s in sheets.values() for m in s["matrics"]]
        self.assertNotIn("2400000005", all_matrics)
        self.assertEqual(len(all_matrics), 5)

    def test_a_slot_with_no_seating_returns_an_explanatory_zip(self):
        SeatArrangement.objects.all().delete()

        response = export_arrangements(RequestFactory().get("/", {"date": DATE, "period": PERIOD}))

        archive = zipfile.ZipFile(io.BytesIO(response.content))
        self.assertIn("error.txt", archive.namelist())


class HallAttendanceApiTests(AttendanceFixture):
    url = reverse("api-export-attendance-sheets")

    def client_for(self, is_staff):
        user = User.objects.create_user(
            email=f"user{is_staff}@example.com", password="x", is_staff=is_staff
        )
        client = APIClient()
        client.force_authenticate(user)
        return client

    def get(self, client, **params):
        query = {"date": DATE, "period": PERIOD, "hall_id": self.hall.id, **params}
        return client.get(self.url, query)

    def test_gives_each_class_its_own_sheet_for_one_hall(self):
        response = self.get(self.client_for(is_staff=True))

        self.assertEqual(response.status_code, 200)
        sheets = read_sheets(response)
        self.assertEqual(len(sheets), 2)
        for sheet in sheets.values():
            self.assertEqual(len({m[-1] in "135" for m in sheet["matrics"]}), 1)
            self.assertEqual(sheet["seats"], sorted(sheet["seats"]))

    def test_sheet_label_matches_the_students_on_it(self):
        sheets = read_sheets(self.get(self.client_for(is_staff=True)))

        for sheet in sheets.values():
            label = "AC HND I" if sheet["matrics"][0].endswith("1") else "BF HND I"
            self.assertIn(f"LEVEL/CLASS: {label}", sheet["text"])

    def test_refuses_a_user_without_an_admin_role(self):
        response = self.get(self.client_for(is_staff=False))

        self.assertEqual(response.status_code, 403)

    def test_refuses_an_anonymous_caller(self):
        response = self.get(APIClient())

        self.assertIn(response.status_code, (401, 403))

    def test_asks_for_missing_parameters(self):
        client = self.client_for(is_staff=True)

        response = client.get(self.url, {"date": DATE, "period": PERIOD})

        self.assertEqual(response.status_code, 400)

    def test_reports_a_hall_with_nobody_seated(self):
        SeatArrangement.objects.update(seat_number=None)

        response = self.get(self.client_for(is_staff=True))

        self.assertEqual(response.status_code, 404)
