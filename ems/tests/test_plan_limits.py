"""Plan limits and features on a school's host (spec 0003, child 0005)."""
import threading
from datetime import timedelta
from unittest import mock

import docx
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection, connections
from django.utils import timezone

from ems.branding import add_document_branding
from ems.models import BackgroundJob, Invite, Student, SystemSettings, User
from ems.tests.helpers import listed_class, make_class, make_department
from ems.tests.test_invites import InviteTestCase
from tenancy import limits
from tenancy.features import KEYS
from tenancy.models import Plan
from tenancy.testcases import SchoolTestCase

PASSWORD = "plum orchard 42"
HEADER = "MATRIC NUMBER,FIRSTNAME,LASTNAME,EMAIL,PHONE NUMBER\n"


def other_connection_can_lock(cap):
    """Whether a second database connection could take this school's
    ``students`` or ``staff`` cap lock right now. A held lock is what makes
    two concurrent writes queue up instead of jointly passing the cap."""
    result = {}
    # Built here: inside the thread, ``connection`` is that thread's own one.
    key = f"{connection.schema_name}:{cap}"
    main = connections["default"]

    def probe():
        other = type(main)(main.settings_dict.copy(), "probe")
        try:
            with other.cursor() as cursor:
                cursor.execute("SELECT pg_try_advisory_lock(hashtext(%s))", [key])
                result["got"] = cursor.fetchone()[0]
                if result["got"]:
                    cursor.execute("SELECT pg_advisory_unlock(hashtext(%s))", [key])
        finally:
            other.close()

    thread = threading.Thread(target=probe)
    thread.start()
    thread.join(10)
    return result["got"]


def students_csv(matrics):
    body = "".join(f"{m},Ada,Obi,{m}@example.com,0803\n" for m in matrics)
    return SimpleUploadedFile("students.csv", (HEADER + body).encode(), content_type="text/csv")


class PlanMixin:
    """Change the test school's plan or overrides for one test.

    Written to the database (requests load the school afresh) and onto the
    in memory school this connection uses, restored after the test.
    """

    def set_school(self, **fields):
        schools = {id(self.tenant): self.tenant, id(connection.tenant): connection.tenant}
        for school in schools.values():
            before = {field: getattr(school, field) for field in fields}
            self.addCleanup(lambda s=school, b=before: [setattr(s, f, v) for f, v in b.items()])
            for field, value in fields.items():
                setattr(school, field, value)
        self.tenant.save(update_fields=list(fields))


def _all_but(features):
    return sorted(KEYS - set(features))


class LimitsTestCase(PlanMixin, SchoolTestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email="sa@test.localhost", password=PASSWORD, role=User.Role.SUPER_ADMIN, is_staff=True,
        )
        self.client.force_authenticate(self.admin)

    def tearDown(self):
        connection.set_tenant(self.tenant)

    def turn_off(self, *features):
        self.set_school(features_override=_all_but(features))


class EffectiveLimitsTests(LimitsTestCase):
    def test_an_override_wins_over_the_plan_and_null_means_unlimited(self):
        # covers: AC-26
        plan = Plan.objects.create(
            code="small", name="Small", max_students=100, max_staff_users=None,
            features=["exports"],
        )
        self.set_school(plan=plan, max_students_override=200)

        effective = limits.effective(self.tenant)

        self.assertEqual(effective.max_students, 200)
        self.assertIsNone(effective.max_staff_users)
        self.assertEqual(effective.features, frozenset({"exports"}))

    def test_an_empty_feature_override_switches_every_feature_off(self):
        self.set_school(features_override=[])

        self.assertEqual(limits.effective(self.tenant).features, frozenset())


class StudentCapTests(LimitsTestCase):
    def setUp(self):
        super().setUp()
        self.cls = make_class(0)

    def upload(self, matrics):
        return self.client.post(
            f"/api/uploads/class-students/{self.cls.pk}/",
            {"file": students_csv(matrics)},
            format="multipart",
        )

    def test_an_upload_over_the_cap_is_refused_whole(self):
        # covers: AC-27
        listed_class(90)
        self.set_school(max_students_override=100)

        response = self.upload([f"NEW{n:03}" for n in range(20)])

        self.assertEqual(response.status_code, 403, response.data)
        self.assertEqual(response.data["code"], "plan_limit")
        self.assertEqual(
            response.data["detail"], "This upload adds 20 students; your plan allows 10 more."
        )
        self.assertEqual(Student.objects.count(), 90)

    def test_rows_that_update_existing_students_dont_count(self):
        # covers: AC-27
        existing = [f"OLD{n:03}" for n in range(50)]
        self.upload(existing)
        listed_class(45)
        self.set_school(max_students_override=100)

        response = self.upload(existing + [f"NEW{n:03}" for n in range(5)])

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data, {"created": 5, "updated": 50})
        self.assertEqual(Student.objects.count(), 100)

    def test_a_plan_override_above_the_plan_lets_a_big_upload_in(self):
        # covers: AC-26
        plan = Plan.objects.create(code="hundred", name="Hundred", max_students=100)
        self.set_school(plan=plan, max_students_override=200)

        response = self.upload([f"NEW{n:03}" for n in range(150)])

        self.assertEqual(response.status_code, 200, response.data)

    def test_adding_one_student_at_the_cap_is_refused(self):
        # covers: AC-27
        listed_class(2)
        self.set_school(max_students_override=2)

        response = self.client.post(
            "/api/students/",
            {"class_id": self.cls.pk, "first_name": "Ada", "last_name": "Obi",
             "matric_no": "NEW001", "email": "ada@example.com"},
            format="json",
        )

        self.assertEqual(response.status_code, 403, response.data)
        self.assertEqual(response.data["detail"], "This adds 1 student; your plan allows 0 more.")
        self.assertEqual(Student.objects.count(), 2)

    def test_lowering_the_cap_below_usage_keeps_existing_students(self):
        listed_class(5)
        self.set_school(max_students_override=3)

        response = self.upload(["NEW001"])

        self.assertEqual(response.status_code, 403)
        self.assertIn("allows 0 more", response.data["detail"])
        self.assertEqual(Student.objects.count(), 5)

    def test_the_cap_must_be_checked_inside_a_transaction(self):
        with self.assertRaises(RuntimeError):
            # TestCase wraps each test in a transaction; leave it for this call.
            with mock.patch.object(connections["default"], "in_atomic_block", False):
                limits.assert_can_add_students(1)


class StaffCapTests(InviteTestCase, PlanMixin):
    def tearDown(self):
        connection.set_tenant(self.tenant)

    def test_an_open_invite_holds_a_seat(self):
        # covers: AC-28
        User.objects.create_user(email="two@test.localhost", password=PASSWORD)
        self.invite(email="three@test.localhost", role="DO")
        self.set_school(max_staff_users_override=3)

        response, _ = self.invite(email="four@test.localhost", role="DO")

        self.assertEqual(response.status_code, 403, response.data)
        self.assertEqual(response.data["code"], "plan_limit")
        self.assertEqual(
            response.data["detail"],
            "Your plan allows 3 staff accounts, and 3 are already in use or invited.",
        )

    def test_an_expired_invite_holds_no_seat(self):
        self.invite(email="three@test.localhost", role="DO")
        Invite.objects.update(expires_at=timezone.now() - timedelta(minutes=1))
        self.set_school(max_staff_users_override=2)

        response, _ = self.invite(email="four@test.localhost", role="DO")

        self.assertEqual(response.status_code, 201, response.data)

    def test_resending_an_invite_doesnt_count_it_twice(self):
        response, _ = self.invite(email="two@test.localhost", role="DO")
        self.set_school(max_staff_users_override=2)

        response, _ = self.invite(path=f"/api/invites/{response.data['id']}/resend/")

        self.assertEqual(response.status_code, 202, response.data)

    def test_accepting_at_the_cap_is_refused(self):
        # covers: AC-28
        _, token = self.invite(email="two@test.localhost", role="DO")
        User.objects.create_user(email="other@test.localhost", password=PASSWORD)
        self.set_school(max_staff_users_override=2)

        response = self.accept(token)

        self.assertEqual(response.status_code, 403, response.data)
        self.assertFalse(User.objects.filter(email="two@test.localhost").exists())

    def test_creating_a_user_at_the_cap_is_refused(self):
        # covers: AC-28
        self.set_school(max_staff_users_override=1)

        response = self.client.post(
            "/api/users/",
            {"email": "new@test.localhost", "password": PASSWORD, "role": "DO"},
            format="json",
        )

        self.assertEqual(response.status_code, 403, response.data)

    def test_an_inactive_user_takes_no_seat_until_reactivated(self):
        # covers: AC-28
        idle = User.objects.create_user(email="idle@test.localhost", password=PASSWORD, is_active=False)
        self.set_school(max_staff_users_override=1)

        response = self.client.patch(f"/api/users/{idle.pk}/", {"is_active": True}, format="json")

        self.assertEqual(response.status_code, 403, response.data)
        idle.refresh_from_db()
        self.assertFalse(idle.is_active)
        # A role change on an existing account isn't a new seat.
        response = self.client.patch(f"/api/users/{self.admin.pk}/", {"role": "SA"}, format="json")
        self.assertEqual(response.status_code, 200, response.data)

    def test_seeding_departments_counts_the_accounts_it_would_create(self):
        # covers: AC-28
        make_department("acc")
        make_department("bus")
        self.set_school(max_staff_users_override=2)

        response = self.client.post("/api/users/seed-departments/", {}, format="json")

        self.assertEqual(response.status_code, 403, response.data)
        self.assertEqual(User.objects.count(), 1)


class FeatureTests(LimitsTestCase):
    def test_seat_allocation_off_gates_its_endpoints(self):
        # covers: AC-29
        self.turn_off("seat_allocation")

        for method, path in [
            ("post", "/api/allocation/generate/"),
            ("post", "/api/allocation/generate-all/"),
            ("post", "/api/allocation/clear/"),
            ("post", "/api/allocation/manual-assign/"),
            ("get", "/api/allocation/"),
            ("get", "/api/allocation/hall/"),
            ("get", "/api/exports/arrangement/"),
            ("get", "/api/exports/attendance-sheets/"),
            ("get", "/api/exports/attendance-sheets/bulk/"),
        ]:
            response = getattr(self.client, method)(path)
            self.assertEqual(response.status_code, 403, path)
            self.assertEqual(response.data["code"], "feature_disabled", path)

    def test_exports_and_publication_off_gate_their_endpoints(self):
        # covers: AC-29
        self.turn_off("exports", "publication")

        for method, path in [
            ("get", "/api/exports/broadsheet/"),
            ("post", "/api/timetable/publish/"),
            ("post", "/api/timetable/unpublish/"),
        ]:
            response = getattr(self.client, method)(path)
            self.assertEqual(response.data.get("code"), "feature_disabled", path)

    def test_a_signed_out_request_is_401_before_any_feature_check(self):
        self.turn_off("seat_allocation")
        self.client.force_authenticate(None)

        self.assertEqual(self.client.post("/api/allocation/generate/").status_code, 401)

    def test_retrying_an_allocation_job_needs_seat_allocation(self):
        # covers: AC-29
        self.turn_off("seat_allocation")
        job = BackgroundJob.objects.create(
            job_id="old", job_type="allocation", status="failed", created_by=self.admin
        )

        response = self.client.post(f"/api/jobs/{job.job_id}/retry/")

        self.assertEqual(response.status_code, 403, response.data)
        self.assertEqual(response.data["code"], "feature_disabled")
        self.assertEqual(BackgroundJob.objects.count(), 1)

    def test_branding_off_refuses_branding_changes_only(self):
        # covers: AC-29
        SystemSettings.objects.create(
            session="2026/2027", semester="1st Semester", exam_heading="Exams Office"
        )
        self.turn_off("branding")

        refused = self.client.patch(
            "/api/system/settings/", {"exam_heading": "New heading"}, format="json"
        )
        self.assertEqual(refused.status_code, 403, refused.data)
        self.assertEqual(refused.data["code"], "feature_disabled")
        # The settings form sends every field; unchanged branding is fine.
        saved = self.client.patch(
            "/api/system/settings/",
            {"exam_heading": "Exams Office", "brand_color": "", "session": "2027/2028"},
            format="json",
        )
        self.assertEqual(saved.status_code, 200, saved.data)
        self.assertEqual(SystemSettings.objects.get().session, "2027/2028")

    def test_branding_off_drops_the_logo_and_heading_from_documents(self):
        # covers: AC-29
        settings_row = SystemSettings.objects.create(
            session="2026/2027", semester="1st Semester", exam_heading="Exams Office",
            institution_name="Test School",
        )
        self.turn_off("branding")
        doc = docx.Document()

        add_document_branding(doc, settings_row, logo=b"not really a png")

        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("EXAMS OFFICE", text)
        self.assertNotIn("[INSTITUTION LOGO]", text)
        self.assertIn("TEST SCHOOL", text)
        self.assertEqual(len(doc.inline_shapes), 0)


class SchoolEndpointTests(LimitsTestCase):
    def test_it_reports_effective_features_limits_and_usage(self):
        # covers: AC-29
        listed_class(3)
        self.set_school(max_students_override=10, features_override=["exports"])

        response = self.client.get("/api/school/")

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["slug"], "test")
        self.assertEqual(response.data["features"], ["exports"])
        self.assertEqual(
            response.data["limits"], {"max_students": 10, "max_staff_users": None}
        )
        self.assertEqual(response.data["usage"], {"students": 3, "staff_users": 1})
        self.assertIsNone(response.data["support"])

    def test_it_needs_a_sign_in(self):
        self.client.force_authenticate(None)

        self.assertEqual(self.client.get("/api/school/").status_code, 401)


class CapLockTests(LimitsTestCase):
    """Concurrent writes: each cap is held until the writing transaction ends,
    so a second upload or accept waits and then counts what the first saved."""

    def test_nothing_holds_the_locks_at_the_start(self):
        self.assertTrue(other_connection_can_lock("students"))
        self.assertTrue(other_connection_can_lock("staff"))

    def test_a_class_upload_holds_the_student_lock_until_it_commits(self):
        # covers: AC-27 (two concurrent uploads can't jointly exceed the cap)
        cls = make_class(0)
        self.set_school(max_students_override=10)

        response = self.client.post(
            f"/api/uploads/class-students/{cls.pk}/",
            {"file": students_csv([f"NEW{n:03}" for n in range(6)])},
            format="multipart",
        )

        self.assertEqual(response.status_code, 200, response.data)
        # The test's own transaction is still open, so the upload's lock is too.
        self.assertFalse(other_connection_can_lock("students"))
        self.assertTrue(other_connection_can_lock("staff"))

    def test_the_lock_is_taken_even_when_nothing_new_is_added(self):
        # The new rows are counted after the lock, so the lock comes first.
        limits.assert_can_add_students(0)

        self.assertFalse(other_connection_can_lock("students"))

    def test_the_lock_is_per_school(self):
        limits.assert_can_add_students(1)

        with mock.patch.object(connection, "schema_name", "school_other"):
            self.assertTrue(other_connection_can_lock("students"))

    def test_creating_a_user_holds_the_staff_lock(self):
        # covers: AC-28 (two concurrent accepts or creates can't both take the last seat)
        response = self.client.post(
            "/api/users/",
            {"email": "new@test.localhost", "password": PASSWORD, "role": "DO"},
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertFalse(other_connection_can_lock("staff"))


class MoreStaffCapTests(LimitsTestCase):
    def test_seeding_within_the_cap_creates_the_accounts(self):
        # covers: AC-28
        make_department("acc")
        self.set_school(max_staff_users_override=2)

        response = self.client.post("/api/users/seed-departments/", {}, format="json")

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["created"], ["acc@ems.com"])

    def test_seeding_skips_existing_accounts_when_counting(self):
        make_department("acc")
        User.objects.create_user(email="acc@ems.com", password=PASSWORD)
        self.set_school(max_staff_users_override=2)

        response = self.client.post("/api/users/seed-departments/", {}, format="json")

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["skipped"], ["acc@ems.com"])

    def test_overwrite_seeding_counts_an_inactive_account_it_turns_back_on(self):
        # covers: AC-28
        make_department("acc")
        User.objects.create_user(email="acc@ems.com", password=PASSWORD, is_active=False)
        self.set_school(max_staff_users_override=1)

        response = self.client.post(
            "/api/users/seed-departments/", {"overwrite": True}, format="json"
        )

        self.assertEqual(response.status_code, 403, response.data)
        self.assertFalse(User.objects.get(email="acc@ems.com").is_active)

    def test_creating_an_inactive_user_takes_no_seat(self):
        self.set_school(max_staff_users_override=1)

        response = self.client.post(
            "/api/users/",
            {"email": "idle@test.localhost", "password": PASSWORD, "role": "DO", "is_active": False},
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.data)

    def test_no_cap_means_no_limit(self):
        for n in range(3):
            User.objects.create_user(email=f"u{n}@test.localhost", password=PASSWORD)

        limits.assert_can_add_staff(100)


class FeatureEdgeTests(LimitsTestCase):
    def test_reading_endpoints_stay_open_with_every_feature_off(self):
        # covers: AC-29 (only the listed endpoints are gated)
        self.set_school(features_override=[])

        for path in [
            "/api/timetable/",
            "/api/readiness/?stage=allocation",
            "/api/distribution/",
            "/api/exports/timetable/",
            "/api/system/branding/",
            "/api/school/",
        ]:
            response = self.client.get(path)
            self.assertNotEqual(response.status_code, 403, path)

    def test_branding_off_still_saves_the_settings_form_without_a_logo(self):
        SystemSettings.objects.create(session="2026/2027", semester="1st Semester")
        self.turn_off("branding")

        response = self.client.patch(
            "/api/system/settings/",
            {"institution_name": "Test School", "exam_heading": "", "brand_color": ""},
            format="multipart",
        )

        self.assertEqual(response.status_code, 200, response.data)

    def test_branding_off_refuses_a_new_brand_colour(self):
        SystemSettings.objects.create(session="2026/2027", semester="1st Semester")
        self.turn_off("branding")

        response = self.client.patch(
            "/api/system/settings/", {"brand_color": "#7C3AED"}, format="json"
        )

        self.assertEqual(response.status_code, 403)
        self.assertEqual(SystemSettings.objects.get().brand_color, "")

    def test_branding_on_lets_the_heading_change(self):
        SystemSettings.objects.create(session="2026/2027", semester="1st Semester")

        response = self.client.patch(
            "/api/system/settings/", {"exam_heading": "Exams Office"}, format="json"
        )

        self.assertEqual(response.status_code, 200, response.data)

    def test_documents_keep_the_heading_when_branding_is_on(self):
        settings_row = SystemSettings.objects.create(
            session="2026/2027", semester="1st Semester", exam_heading="Exams Office"
        )
        doc = docx.Document()

        add_document_branding(doc, settings_row, logo=None)

        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("EXAMS OFFICE", text)
        self.assertIn("[INSTITUTION LOGO]", text)

    def test_a_feature_message_names_the_feature(self):
        self.turn_off("publication")

        with self.assertRaisesMessage(
            limits.FeatureDisabled, "Timetable publishing isn't included in your plan."
        ):
            limits.assert_feature("publication")

    def test_retrying_a_timetable_job_needs_no_feature(self):
        self.set_school(features_override=[])

        limits.assert_feature_for_job("timetable")
        limits.assert_feature_for_job("distribution_all")
        with self.assertRaises(limits.FeatureDisabled):
            limits.assert_feature_for_job("allocation_all")


class ErrorEnvelopeTests(LimitsTestCase):
    def test_other_errors_carry_no_code(self):
        # Only plan errors add ``code``; the rest of the envelope is unchanged.
        response = self.client.delete(f"/api/users/{self.admin.pk}/")

        self.assertEqual(response.status_code, 409)
        self.assertEqual(set(response.data), {"detail"})
